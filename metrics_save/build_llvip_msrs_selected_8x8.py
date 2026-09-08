#!/usr/bin/env python3
"""Build filtered LLVIP+MSRS Excel: 8 methods x 8 metrics (Fusion_training_1 optimized set)."""

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

PROJECT = Path('/root/autodl-tmp/URFusion-main')
METRICS_SAVE = PROJECT / 'metrics_save'

METHODS = [
	'Fusion_training_1',
	'MUFusion',
	'U2Fusion',
	'URFusion',
	'MetaFusion',
	'GIFNet',
	'SAGE',
	'LRRNet',
]
METRICS = ['NMI', 'Qy', 'EN', 'MI', 'VIF', 'Qabf', 'SSIM', 'VIFF']
DATASETS = ['LLVIP', 'MSRS']

OUT_XLSX = METRICS_SAVE / 'LLVIP_MSRS_selected_8x8.xlsx'
OUT_COLORED = METRICS_SAVE / 'LLVIP_MSRS_selected_8x8_colored.xlsx'


def load_source_rows(dataset: str):
	wb = openpyxl.load_workbook(METRICS_SAVE / f'{dataset}.xlsx', data_only=True)
	ws = wb['MeanMetrics']
	headers = [c.value for c in ws[1]]
	rows = {}
	for row in ws.iter_rows(min_row=2, values_only=True):
		if not row[0]:
			continue
		rows[row[0]] = dict(zip(headers, row))
	return rows


def build_sheet(ws, dataset: str, source_rows: dict):
	ws.append(['Method'] + METRICS)
	for method in METHODS:
		if method not in source_rows:
			raise KeyError(f'{method} missing in {dataset}.xlsx')
		row = source_rows[method]
		ws.append([method] + [row[m] for m in METRICS])


BEST_COLOR = 'FFFF0000'
SECOND_COLOR = 'FF0000FF'
HIGHER_BETTER = {
	'NMI', 'Qy', 'EN', 'MI', 'VIF', 'Qabf', 'SSIM', 'VIFF',
	'QNCIE', 'EI', 'Qcb', 'SF', 'AG', 'SD', 'CC', 'SCD', 'PSNR', 'MS_SSIM', 'MUSIQ',
}
LOWER_BETTER = {'CE', 'TE', 'MSE', 'Nabf', 'NIQE', 'BRISQUE'}


def to_float(value):
	try:
		return float(value)
	except (TypeError, ValueError):
		return None


def recolor_font(cell, rgb: str):
	cell.font = Font(color=rgb, bold=cell.font.bold if cell.font else False)


def apply_ranking_colors(ws):
	headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
	if not headers or len(headers) < 2:
		raise ValueError('表头不完整')

	for col_idx in range(2, len(headers) + 1):
		metric = headers[col_idx - 1]
		if not metric:
			continue
		if metric in LOWER_BETTER:
			reverse = False
		elif metric in HIGHER_BETTER:
			reverse = True
		else:
			continue

		numeric_items = []
		for row_idx in range(2, ws.max_row + 1):
			val = to_float(ws.cell(row_idx, col_idx).value)
			if val is not None:
				numeric_items.append((row_idx, val))
		if not numeric_items:
			continue

		unique_values = sorted({val for _, val in numeric_items}, reverse=reverse)
		best_value = unique_values[0]
		second_value = unique_values[1] if len(unique_values) > 1 else None

		for row_idx, val in numeric_items:
			cell = ws.cell(row_idx, col_idx)
			if val == best_value:
				recolor_font(cell, BEST_COLOR)
			elif second_value is not None and val == second_value:
				recolor_font(cell, SECOND_COLOR)


def main():
	plain = openpyxl.Workbook()
	plain.remove(plain.active)
	colored = openpyxl.Workbook()
	colored.remove(colored.active)

	for dataset in DATASETS:
		source_rows = load_source_rows(dataset)
		pws = plain.create_sheet(title=dataset)
		build_sheet(pws, dataset, source_rows)

		cws = colored.create_sheet(title=dataset)
		build_sheet(cws, dataset, source_rows)
		apply_ranking_colors(cws)

	METRICS_SAVE.mkdir(parents=True, exist_ok=True)
	plain.save(OUT_XLSX)
	colored.save(OUT_COLORED)
	print(f'Wrote {OUT_XLSX}')
	print(f'Wrote {OUT_COLORED}')
	print(f'Methods ({len(METHODS)}): {", ".join(METHODS)}')
	print(f'Metrics ({len(METRICS)}): {", ".join(METRICS)}')


if __name__ == '__main__':
	main()
