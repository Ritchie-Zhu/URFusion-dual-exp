#!/usr/bin/env python3
"""Parallel 25-metric eval using the official eval_torch.evaluation_one bytecode.

Same Excel / json / csv layout as eval_multi_time. Does not modify metrics_py/*.py.
"""
from __future__ import annotations

import argparse
import csv
import importlib.machinery
import importlib.util
import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import numpy as np
from natsort import natsorted
from openpyxl import Workbook, load_workbook
from tqdm import tqdm

PROJECT = Path('/root/autodl-tmp/URFusion-main')
METRICS_PY = str(PROJECT / 'metrics_py')
IMAGE_EXTS = ('.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff')
STATIC_METRIC_NAMES = [
	'CE', 'NMI', 'QNCIE', 'TE', 'EI', 'Qy', 'Qcb', 'EN', 'MI', 'SF', 'AG', 'SD',
	'CC', 'SCD', 'VIF', 'MSE', 'PSNR', 'Qabf', 'Nabf', 'SSIM', 'MS_SSIM', 'VIFF',
	'NIQE', 'BRISQUE', 'MUSIQ',
]
SHEET_NAME = 'MeanMetrics'

_et = None


def load_pyc(name: str):
	pyc = os.path.join(METRICS_PY, '__pycache__', f'{name}.cpython-310.pyc')
	if not os.path.isfile(pyc):
		raise FileNotFoundError(f'Missing bytecode cache: {pyc}')
	loader = importlib.machinery.SourcelessFileLoader(name, pyc)
	spec = importlib.util.spec_from_loader(name, loader)
	mod = importlib.util.module_from_spec(spec)
	mod.__file__ = os.path.join(METRICS_PY, f'{name}.py')
	sys.modules[name] = mod
	loader.exec_module(mod)
	return mod


def to_scalar(value):
	try:
		import torch

		if isinstance(value, torch.Tensor):
			if value.numel() == 1:
				return float(value.detach().cpu().item())
			return float(value.detach().cpu().mean().item())
	except Exception:
		pass
	if isinstance(value, np.generic):
		return float(value.item())
	return float(value)


def init_worker():
	global _et
	os.environ['OMP_NUM_THREADS'] = '1'
	os.environ['MKL_NUM_THREADS'] = '1'
	os.environ['OPENBLAS_NUM_THREADS'] = '1'
	os.environ['NUMEXPR_NUM_THREADS'] = '1'
	os.chdir(METRICS_PY)
	if METRICS_PY not in sys.path:
		sys.path.insert(0, METRICS_PY)
	try:
		import torch

		torch.set_num_threads(1)
	except Exception:
		pass
	load_pyc('Metric_torch')
	_et = load_pyc('eval_torch')


def eval_one(task):
	name, ir_path, vis_path, fused_path = task
	t0 = time.perf_counter()
	result = _et.evaluation_one(ir_path, vis_path, fused_path)
	elapsed = time.perf_counter() - t0
	row = {'image_name': name, 'time_s': elapsed}
	for key, value in zip(_et.STATIC_METRIC_NAMES, result):
		row[key] = to_scalar(value)
	return row


def parse_args():
	parser = argparse.ArgumentParser(description='Parallel static fusion metrics (official evaluation_one)')
	parser.add_argument('--methods', nargs='+', required=True)
	parser.add_argument('--datasets', nargs='+', required=True)
	parser.add_argument('--dataset_roots', nargs='+', required=True)
	parser.add_argument('--results_root', type=str, default=str(PROJECT / 'our_model_1_DualMoE/results'))
	parser.add_argument('--metrics_save', type=str, default=str(PROJECT / 'metrics_save'))
	parser.add_argument('--project_root', type=str, default=str(PROJECT))
	parser.add_argument('--num_workers', type=int, default=8)
	parser.add_argument('--force', action='store_true', help='Recompute even if metrics_summary.json exists')
	return parser.parse_args()


def resolve_vis_ir_dirs(test_dir):
	test_dir = os.path.abspath(test_dir)
	vis_dir = None
	for name in ('VIS', 'vis'):
		candidate = os.path.join(test_dir, name)
		if os.path.isdir(candidate):
			vis_dir = candidate
			break
	if vis_dir is None:
		raise FileNotFoundError(f'VIS/vis not found under {test_dir}')
	parent = os.path.dirname(vis_dir.rstrip(os.sep))
	ir_dir = None
	for name in ('IR', 'ir'):
		candidate = os.path.join(parent, name)
		if os.path.isdir(candidate):
			ir_dir = candidate
			break
	if ir_dir is None:
		raise FileNotFoundError(f'IR/ir not found next to {vis_dir}')
	return vis_dir, ir_dir


def resolve_source_name(vis_dir, ir_dir, fused_name):
	if os.path.isfile(os.path.join(vis_dir, fused_name)) and os.path.isfile(os.path.join(ir_dir, fused_name)):
		return fused_name
	stem, _ = os.path.splitext(fused_name)
	for ext in IMAGE_EXTS:
		candidate = stem + ext
		if os.path.isfile(os.path.join(vis_dir, candidate)) and os.path.isfile(os.path.join(ir_dir, candidate)):
			return candidate
	return None


def list_paired_names(vis_dir, ir_dir, fused_dir):
	pairs = []
	for fused_fn in natsorted(os.listdir(fused_dir)):
		if os.path.splitext(fused_fn)[1].lower() not in IMAGE_EXTS:
			continue
		source_name = resolve_source_name(vis_dir, ir_dir, fused_fn)
		if source_name:
			pairs.append((source_name, fused_fn))
	return pairs


def prepare_sheet_for_append(excel_path, sheet_name, header):
	try:
		wb = load_workbook(excel_path)
	except FileNotFoundError:
		wb = Workbook()
		ws = wb.active
		ws.title = sheet_name
	else:
		if sheet_name in wb.sheetnames:
			ws = wb[sheet_name]
		else:
			ws = wb.create_sheet(sheet_name)
	for col, value in enumerate(header, start=1):
		ws.cell(row=1, column=col, value=value)
	wb.save(excel_path)


def remove_method_rows(excel_path, sheet_name, method):
	try:
		wb = load_workbook(excel_path)
	except FileNotFoundError:
		return
	if sheet_name not in wb.sheetnames:
		return
	ws = wb[sheet_name]
	removed = 0
	for row_idx in range(ws.max_row, 1, -1):
		if ws.cell(row_idx, 1).value == method:
			ws.delete_rows(row_idx)
			removed += 1
	if removed:
		wb.save(excel_path)


def append_to_excel(excel_path, sheet_name, new_row):
	try:
		wb = load_workbook(excel_path)
	except FileNotFoundError:
		wb = Workbook()
	if sheet_name not in wb.sheetnames:
		wb.create_sheet(sheet_name)
	ws = wb[sheet_name]
	next_row = ws.max_row + 1
	for col, val in enumerate(new_row, start=1):
		ws.cell(row=next_row, column=col, value=val)
	wb.save(excel_path)


def xlsx_has_method(excel_path, method):
	try:
		wb = load_workbook(excel_path, data_only=True)
	except FileNotFoundError:
		return False
	if SHEET_NAME not in wb.sheetnames:
		return False
	ws = wb[SHEET_NAME]
	for row_idx in range(2, ws.max_row + 1):
		if ws.cell(row_idx, 1).value == method:
			return True
	return False


def save_method_outputs(save_dir, method, dataset, summary, per_image_rows):
	tag_dir = os.path.join(save_dir, f'{method}_{dataset}')
	os.makedirs(tag_dir, exist_ok=True)
	summary_path = os.path.join(tag_dir, 'metrics_summary.json')
	payload = {'method': method, 'dataset': dataset, **summary}
	with open(summary_path, 'w', encoding='utf-8') as f:
		json.dump(payload, f, indent=2, sort_keys=True)
	csv_path = os.path.join(tag_dir, 'metrics_per_image.csv')
	fields = ['image_name', 'time_s'] + STATIC_METRIC_NAMES
	with open(csv_path, 'w', newline='', encoding='utf-8') as f:
		writer = csv.DictWriter(f, fieldnames=fields)
		writer.writeheader()
		for row in per_image_rows:
			writer.writerow(row)
	return summary_path, csv_path


def load_summary(path):
	with open(path, encoding='utf-8') as f:
		return json.load(f)


def summarize(per_image_rows):
	all_values = {k: [] for k in STATIC_METRIC_NAMES}
	times = []
	for row in per_image_rows:
		times.append(float(row['time_s']))
		for metric_name in STATIC_METRIC_NAMES:
			all_values[metric_name].append(float(row[metric_name]))
	summary = {}
	for metric_name in STATIC_METRIC_NAMES:
		arr = np.array(all_values[metric_name], dtype=np.float64)
		summary[f'{metric_name}_mean'] = float(arr.mean())
		summary[f'{metric_name}_std'] = float(arr.std(ddof=0))
	arr_t = np.array(times, dtype=np.float64)
	summary['n_images'] = len(per_image_rows)
	summary['time_per_image_mean_s'] = float(arr_t.mean())
	summary['time_per_image_std_s'] = float(arr_t.std(ddof=0))
	return summary


def excel_row_from_summary(method, summary):
	row = [method]
	for metric_name in STATIC_METRIC_NAMES:
		val = summary.get(f'{metric_name}_mean', float('nan'))
		row.append(round(float(val), 4) if np.isfinite(val) else None)
	return row


def append_summary_to_excel(excel_path, method, summary):
	remove_method_rows(excel_path, SHEET_NAME, method)
	append_to_excel(excel_path, SHEET_NAME, excel_row_from_summary(method, summary))


def evaluate_with_pool(pool, method, dataset, test_dir, fused_dir):
	vis_dir, ir_dir = resolve_vis_ir_dirs(test_dir)
	pairs = list_paired_names(vis_dir, ir_dir, fused_dir)
	if not pairs:
		raise FileNotFoundError(f'No paired images for {method}_{dataset} in {fused_dir}')
	tasks = [
		(
			source_name,
			os.path.join(ir_dir, source_name),
			os.path.join(vis_dir, source_name),
			os.path.join(fused_dir, fused_fn),
		)
		for source_name, fused_fn in pairs
	]
	per_image_rows = []
	futures = {pool.submit(eval_one, task): task[0] for task in tasks}
	for fut in tqdm(as_completed(futures), total=len(futures), desc=f'{method}_{dataset}', leave=True):
		per_image_rows.append(fut.result())
	per_image_rows.sort(key=lambda r: r['image_name'])
	return summarize(per_image_rows), per_image_rows


def main():
	from multiprocessing import set_start_method

	set_start_method('spawn', force=True)
	args = parse_args()
	if len(args.dataset_roots) != len(args.datasets):
		raise ValueError('--dataset_roots length must match --datasets')

	os.makedirs(args.metrics_save, exist_ok=True)
	dataset_roots = dict(zip(args.datasets, args.dataset_roots))
	header = ['Method'] + STATIC_METRIC_NAMES
	for dataset in args.datasets:
		excel_path = os.path.join(args.metrics_save, f'{dataset}.xlsx')
		prepare_sheet_for_append(excel_path, SHEET_NAME, header)

	print(f'[eval_mp] methods={args.methods}')
	print(f'[eval_mp] datasets={args.datasets}')
	print(f'[eval_mp] num_workers={args.num_workers}')
	print(f'[eval_mp] results_root={os.path.abspath(args.results_root)}')

	need_compute = []
	for method in args.methods:
		for dataset in args.datasets:
			fused_dir = os.path.join(args.results_root, f'{method}_{dataset}', 'RGB_fused')
			excel_path = os.path.join(args.metrics_save, f'{dataset}.xlsx')
			summary_path = os.path.join(args.metrics_save, f'{method}_{dataset}', 'metrics_summary.json')
			test_dir = dataset_roots[dataset]
			if not os.path.isdir(fused_dir):
				raise FileNotFoundError(f'fused dir missing: {fused_dir}')
			vis_dir, ir_dir = resolve_vis_ir_dirs(test_dir)
			n_pairs = len(list_paired_names(vis_dir, ir_dir, fused_dir))
			if (
				not args.force
				and os.path.isfile(summary_path)
				and load_summary(summary_path).get('n_images') == n_pairs
			):
				summary = load_summary(summary_path)
				if not xlsx_has_method(excel_path, method):
					append_summary_to_excel(excel_path, method, summary)
					print(f'[eval_mp] restored Excel row {method} @ {dataset} from existing json')
				else:
					print(f'[eval_mp] skip {method} @ {dataset} (existing n={n_pairs})')
				continue
			need_compute.append((method, dataset, test_dir, fused_dir, n_pairs))

	if need_compute:
		with ProcessPoolExecutor(max_workers=args.num_workers, initializer=init_worker) as pool:
			for method, dataset, test_dir, fused_dir, n_pairs in need_compute:
				excel_path = os.path.join(args.metrics_save, f'{dataset}.xlsx')
				print(f'===== {method} @ {dataset} =====')
				t0 = time.time()
				summary, per_image_rows = evaluate_with_pool(pool, method, dataset, test_dir, fused_dir)
				if summary['n_images'] != n_pairs:
					raise RuntimeError(f'{method}_{dataset}: got {summary["n_images"]} expected {n_pairs}')
				summary_path, csv_path = save_method_outputs(
					args.metrics_save, method, dataset, summary, per_image_rows
				)
				append_summary_to_excel(excel_path, method, summary)
				elapsed = time.time() - t0
				print(
					f'  n={summary["n_images"]} elapsed={elapsed:.1f}s '
					f'单张平均={elapsed / summary["n_images"]:.3f}s/img'
				)
				print(f'  summary -> {summary_path}')
				print(f'  per-image -> {csv_path}')
				print(f'  excel -> {excel_path}')

	print(f'All done. metrics_save={os.path.abspath(args.metrics_save)}')


if __name__ == '__main__':
	main()
