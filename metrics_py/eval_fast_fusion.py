"""
Fast Tier-A metrics on RGB_fused outputs (10 metrics, multiprocessing).

Usage:
  python eval_fast_fusion.py --methods Fusion_training_1 --datasets M3FD --num_workers 8
"""

import argparse
import csv
import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed

import numpy as np
from natsort import natsorted
from openpyxl import Workbook, load_workbook
from tqdm import tqdm

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)

from eval_torch_fast import FAST_METRIC_NAMES, evaluation_one_fast

IMAGE_EXTS = ('.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff')


def parse_args():
	parser = argparse.ArgumentParser(description='Fast Tier-A fusion metrics (RGB_fused)')
	parser.add_argument('--methods', nargs='+', default=['Fusion_training_1'])
	parser.add_argument('--datasets', nargs='+', default=['M3FD'])
	parser.add_argument('--dataset_roots', nargs='+', default=[])
	parser.add_argument(
		'--results_root',
		type=str,
		default='/root/autodl-tmp/URFusion-main/vis-ir-gray/results',
	)
	parser.add_argument(
		'--metrics_save',
		type=str,
		default='/root/autodl-tmp/URFusion-main/metrics_save',
	)
	parser.add_argument('--project_root', type=str, default='/root/autodl-tmp/URFusion-main')
	parser.add_argument('--num_workers', type=int, default=8)
	return parser.parse_args()


def resolve_vis_ir_dirs(test_dir):
	test_dir = os.path.abspath(test_dir)
	vis_dir = ir_dir = None
	for name in ('VIS', 'vis'):
		candidate = os.path.join(test_dir, name)
		if os.path.isdir(candidate):
			vis_dir = candidate
			break
	if vis_dir is None:
		raise FileNotFoundError(f'VIS/vis not found under {test_dir}')
	parent = os.path.dirname(vis_dir.rstrip(os.sep))
	for name in ('IR', 'ir'):
		candidate = os.path.join(parent, name)
		if os.path.isdir(candidate):
			ir_dir = candidate
			break
	if ir_dir is None:
		raise FileNotFoundError(f'IR/ir not found next to {vis_dir}')
	return vis_dir, ir_dir


def list_paired_names(vis_dir, ir_dir, fused_dir):
	names = []
	for fn in natsorted(os.listdir(fused_dir)):
		if os.path.splitext(fn)[1].lower() not in IMAGE_EXTS:
			continue
		if os.path.isfile(os.path.join(vis_dir, fn)) and os.path.isfile(os.path.join(ir_dir, fn)):
			names.append(fn)
	return names


def prepare_sheet_for_append(excel_path, sheet_name, header):
	try:
		wb = load_workbook(excel_path)
	except FileNotFoundError:
		wb = Workbook()
		ws = wb.active
		ws.title = sheet_name
	else:
		if sheet_name in wb.sheetnames:
			ws = wb[sheet_name]
		else:
			ws = wb.create_sheet(sheet_name)
	for col, value in enumerate(header, start=1):
		ws.cell(row=1, column=col, value=value)
	wb.save(excel_path)


def append_to_excel(excel_path, sheet_name, new_row):
	try:
		wb = load_workbook(excel_path)
	except FileNotFoundError:
		wb = Workbook()
	if sheet_name not in wb.sheetnames:
		wb.create_sheet(sheet_name)
	ws = wb[sheet_name]
	next_row = ws.max_row + 1
	for col, val in enumerate(new_row, start=1):
		ws.cell(row=next_row, column=col, value=val)
	wb.save(excel_path)


def resolve_dataset_roots(args):
	if args.dataset_roots:
		if len(args.dataset_roots) != len(args.datasets):
			raise ValueError('--dataset_roots length must match --datasets')
		return dict(zip(args.datasets, args.dataset_roots))
	return {
		name: os.path.join(args.project_root, 'datasets', name, 'test')
		for name in args.datasets
	}


def _eval_one_triplet(paths):
	name, ir_path, vi_path, f_path = paths
	metrics = evaluation_one_fast(ir_path, vi_path, f_path)
	row = {'image_name': name}
	row.update(metrics)
	return row


def evaluate_method_dataset(method, dataset, test_dir, fused_dir, num_workers):
	vis_dir, ir_dir = resolve_vis_ir_dirs(test_dir)
	names = list_paired_names(vis_dir, ir_dir, fused_dir)
	if not names:
		raise FileNotFoundError(f'No paired images for {method}_{dataset} in {fused_dir}')

	tasks = [
		(name, os.path.join(ir_dir, name), os.path.join(vis_dir, name), os.path.join(fused_dir, name))
		for name in names
	]

	per_image_rows = []
	if num_workers <= 1:
		for task in tqdm(tasks, desc=f'{method}_{dataset}', leave=False):
			per_image_rows.append(_eval_one_triplet(task))
	else:
		with ProcessPoolExecutor(max_workers=num_workers) as pool:
			futures = {pool.submit(_eval_one_triplet, t): t[0] for t in tasks}
			for fut in tqdm(as_completed(futures), total=len(futures), desc=f'{method}_{dataset}', leave=False):
				per_image_rows.append(fut.result())

	per_image_rows.sort(key=lambda r: r['image_name'])

	all_values = {k: [] for k in FAST_METRIC_NAMES}
	for row in per_image_rows:
		for metric_name in FAST_METRIC_NAMES:
			all_values[metric_name].append(float(row[metric_name]))

	summary = {'profile': 'fast', 'metrics': FAST_METRIC_NAMES}
	for metric_name in FAST_METRIC_NAMES:
		arr = np.array(all_values[metric_name], dtype=np.float64)
		summary[f'{metric_name}_mean'] = float(arr.mean())
		summary[f'{metric_name}_std'] = float(arr.std(ddof=0))
	summary['n_images'] = len(names)
	return summary, per_image_rows


def save_method_outputs(save_dir, method, dataset, summary, per_image_rows):
	tag_dir = os.path.join(save_dir, f'{method}_{dataset}')
	os.makedirs(tag_dir, exist_ok=True)

	summary_path = os.path.join(tag_dir, 'metrics_summary_fast.json')
	payload = {'method': method, 'dataset': dataset, **summary}
	with open(summary_path, 'w', encoding='utf-8') as f:
		json.dump(payload, f, indent=2, sort_keys=True)

	csv_path = os.path.join(tag_dir, 'metrics_per_image_fast.csv')
	with open(csv_path, 'w', newline='', encoding='utf-8') as f:
		fields = ['image_name'] + FAST_METRIC_NAMES
		writer = csv.DictWriter(f, fieldnames=fields)
		writer.writeheader()
		for row in per_image_rows:
			writer.writerow(row)
	return summary_path, csv_path


def main():
	args = parse_args()
	os.makedirs(args.metrics_save, exist_ok=True)
	dataset_roots = resolve_dataset_roots(args)
	sheet_name = 'FastMetrics'
	header = ['Method'] + FAST_METRIC_NAMES

	for dataset in args.datasets:
		excel_path = os.path.join(args.metrics_save, f'{dataset}_fast.xlsx')
		prepare_sheet_for_append(excel_path, sheet_name, header)

	print(f'[fast] metrics={FAST_METRIC_NAMES}')
	print(f'[fast] num_workers={args.num_workers}')

	for method in args.methods:
		for dataset in args.datasets:
			test_dir = dataset_roots[dataset]
			fused_dir = os.path.join(args.results_root, f'{method}_{dataset}', 'RGB_fused')
			if not os.path.isdir(fused_dir):
				print(f'[SKIP] fused dir missing: {fused_dir}')
				continue

			print(f'===== {method} @ {dataset} =====')
			t0 = time.time()
			summary, per_image_rows = evaluate_method_dataset(
				method, dataset, test_dir, fused_dir, num_workers=args.num_workers
			)
			summary_path, csv_path = save_method_outputs(
				args.metrics_save, method, dataset, summary, per_image_rows
			)

			excel_path = os.path.join(args.metrics_save, f'{dataset}_fast.xlsx')
			row = [method]
			for metric_name in FAST_METRIC_NAMES:
				val = summary.get(f'{metric_name}_mean', float('nan'))
				row.append(round(float(val), 4) if np.isfinite(val) else None)
			append_to_excel(excel_path, sheet_name, row)

			elapsed = time.time() - t0
			print(f'  n={summary["n_images"]} elapsed={elapsed:.1f}s ({elapsed / summary["n_images"]:.2f}s/img)')
			for metric_name in FAST_METRIC_NAMES:
				print(f'    {metric_name}: {summary[f"{metric_name}_mean"]:.4f}')
			print(f'  summary -> {summary_path}')
			print(f'  per-image -> {csv_path}')
			print(f'  excel -> {excel_path}')

	print(f'All done. metrics_save={os.path.abspath(args.metrics_save)}')


if __name__ == '__main__':
	main()
