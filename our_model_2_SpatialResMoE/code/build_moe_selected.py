#!/usr/bin/env python3
"""Build SpatialResMoE_3 selected SOTA table (exclude Full / A2).

Same rule as metrics_save/Final_selected.xlsx:
among the 7 comparison methods, keep a metric only if SpatialResMoE_3
is strictly better on LLVIP, MSRS, and M3FD. That set is still
NMI, Qy, MI, VIF, Qabf, VIFF.
"""

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

PROJECT = Path('/root/autodl-tmp/URFusion-main')
OURS = 'SpatialResMoE_3'
OURS_ROOT = PROJECT / 'our_model_2_SpatialResMoE' / 'eval_save'
METRICS_SAVE = PROJECT / 'metrics_save'

METHODS = [
	OURS,
	'MUFusion',
	'U2Fusion',
	'URFusion',
	'MetaFusion',
	'GIFNet',
	'SAGE',
	'LRRNet',
]
METRICS = ['NMI', 'Qy', 'MI', 'VIF', 'Qabf', 'VIFF']
DATASETS = ['LLVIP', 'MSRS', 'M3FD']
OUT_XLSX = OURS_ROOT / 'SpatialResMoE_3_selected.xlsx'
OUT_COLORED = OURS_ROOT / 'SpatialResMoE_3_selected_colored.xlsx'
BEST_COLOR = 'FFFF0000'
SECOND_COLOR = 'FF0000FF'


def load_row(method: str, dataset: str):
	root = OURS_ROOT if method == OURS else METRICS_SAVE
	path = root / f'{method}_{dataset}' / 'metrics_summary.json'
	with open(path, encoding='utf-8') as f:
		data = json.load(f)
	return {k.replace('_mean', ''): data[k] for k in data if k.endswith('_mean')}


def build_sheet(ws, dataset: str):
	ws.append(['Method'] + METRICS)
	for method in METHODS:
		row = load_row(method, dataset)
		ws.append([method] + [round(float(row[m]), 4) for m in METRICS])


def apply_ranking_colors(ws):
	for col_idx in range(2, len(METRICS) + 2):
		items = []
		for row_idx in range(2, ws.max_row + 1):
			val = ws.cell(row_idx, col_idx).value
			if isinstance(val, (int, float)):
				items.append((row_idx, float(val)))
		unique = sorted({v for _, v in items}, reverse=True)
		best = unique[0]
		second = unique[1] if len(unique) > 1 else None
		for row_idx, val in items:
			cell = ws.cell(row_idx, col_idx)
			if val == best:
				cell.font = Font(color=BEST_COLOR, bold=True)
			elif second is not None and val == second:
				cell.font = Font(color=SECOND_COLOR)


def main():
	plain = openpyxl.Workbook()
	plain.remove(plain.active)
	colored = openpyxl.Workbook()
	colored.remove(colored.active)
	for dataset in DATASETS:
		pws = plain.create_sheet(title=dataset)
		build_sheet(pws, dataset)
		cws = colored.create_sheet(title=dataset)
		build_sheet(cws, dataset)
		apply_ranking_colors(cws)
	OURS_ROOT.mkdir(parents=True, exist_ok=True)
	plain.save(OUT_XLSX)
	colored.save(OUT_COLORED)
	print(f'Wrote {OUT_XLSX}')
	print(f'Wrote {OUT_COLORED}')


if __name__ == '__main__':
	main()
